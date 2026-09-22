"""Re-open the surgically edited fixtures with the REFERENCE libraries.

The vitest suite proves the writer is structurally correct and that untouched
archive entries are byte-identical. This script proves the stronger claim that
actually matters: an independent, real implementation opens the result and finds
the formatting, formulas and layouts intact.

Run AFTER the vitest suite has produced tests/fixtures/ooxml/out/:
    node node_modules/vitest/vitest.mjs run tests/frontend/lib/ooxml/real-fixtures.test.ts
    python tests/fixtures/ooxml/verify_roundtrip.py
"""

from __future__ import annotations

import pathlib
import sys

from openpyxl import load_workbook
from pptx import Presentation

HERE = pathlib.Path(__file__).resolve().parent
OUT = HERE / "out"

# The em dash in the edited title is not representable in the default Windows
# console code page; without this the output reads "Q4 Review ?? edited" and
# looks like corruption when the file is in fact correct.
if hasattr(sys.stdout, "reconfigure"):
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")

failures: list[str] = []


def check(label: str, actual: object, expected: object) -> None:
    if actual == expected:
        print(f"  ok   {label} = {actual!r}")
    else:
        failures.append(f"{label}: expected {expected!r}, got {actual!r}")
        print(f"  FAIL {label}: expected {expected!r}, got {actual!r}")


def verify_xlsx() -> None:
    path = OUT / "sample-edited.xlsx"
    print(f"\n== openpyxl re-opens {path.name} ==")
    book = load_workbook(path)

    check("sheet names", book.sheetnames, ["Report", "Notes"])
    report = book["Report"]

    # --- the edits landed ---------------------------------------------------
    check("B3 edited value", report["B3"].value, 999)
    check("C3 edited value", report["C3"].value, 19980)
    check("A9 added cell", report["A9"].value, "Added by the editor")

    # --- everything the editor never touched survived -----------------------
    check("A1 title text", report["A1"].value, "Quarterly Report")
    check("A1 merge range", [str(r) for r in report.merged_cells.ranges], ["A1:C1"])
    check("A1 bold", report["A1"].font.bold, True)
    check("A2 header bold", report["A2"].font.bold, True)
    check("A2 header fill", report["A2"].fill.fgColor.rgb, "00B1552F")
    check("C4 number format", report["C4"].number_format, "#,##0.00")
    check("B6 formula text", report["B6"].value, "=SUM(B3:B5)")
    check("B6 bold", report["B6"].font.bold, True)
    check("column A width", report.column_dimensions["A"].width, 24)
    check("freeze panes", report.freeze_panes, "A3")
    check("untouched B4", report["B4"].value, 80)
    check("untouched A3", report["A3"].value, "Widgets")

    # --- the second sheet is a different part and must be pristine ----------
    notes = book["Notes"]
    check("Notes A1", notes["A1"].value, "Shared string alpha")
    check("Notes B1 boolean", notes["B1"].value, True)
    check("Notes B2 percent format", notes["B2"].number_format, "0.0%")


def verify_pptx() -> None:
    path = OUT / "sample-edited.pptx"
    print(f"\n== python-pptx re-opens {path.name} ==")
    deck = Presentation(path)

    check("slide count", len(deck.slides), 2)
    check("slide 1 title", deck.slides[0].shapes.title.text, "Q4 Review — edited")
    check(
        "slide 1 subtitle preserved",
        deck.slides[0].placeholders[1].text,
        "Prepared by the data team",
    )
    check("slide 2 title preserved", deck.slides[1].shapes.title.text, "Highlights")
    check(
        "slide 2 bullets preserved",
        deck.slides[1].placeholders[1].text_frame.text,
        "Revenue up 12%\nChurn down to 1.4%",
    )
    check("layouts preserved", len(deck.slide_layouts), 11)


def verify_browser_outputs() -> None:
    """Open files produced by the editors IN REAL CHROME.

    This is the end-to-end claim: the browser edited the document and the
    reference implementation still accepts it. Only produced when
    `node tests/browser/run-editors.mjs` has been run first.
    """
    print("\n== reference libraries open the BROWSER-produced files ==")

    sheet_path = OUT / "browser-sheet.xlsx"
    if sheet_path.exists():
        book = load_workbook(sheet_path)
        report = book["Report"]
        check("browser xlsx: edited cell", report["A3"].value, "WidgetsPlus")
        check("browser xlsx: other cells intact", report["B4"].value, 80)
        check("browser xlsx: formula intact", report["B6"].value, "=SUM(B3:B5)")
        check("browser xlsx: styles intact", report["A2"].fill.fgColor.rgb, "00B1552F")
        check("browser xlsx: second sheet intact", book["Notes"]["A1"].value, "Shared string alpha")
    else:
        print(f"  skip {sheet_path.name} (run the browser suite first)")

    docx_path = OUT / "browser-docx.docx"
    if docx_path.exists():
        from docx import Document

        document = Document(docx_path)
        text = "\n".join(p.text for p in document.paragraphs)
        check("browser docx: heading intact", "Quarterly Report" in text, True)
        check("browser docx: typed text present", "Edited in Chrome" in text, True)
        check("browser docx: table intact", document.tables[0].cell(1, 0).text, "Widgets")
    else:
        print(f"  skip {docx_path.name} (run the browser suite first)")

    pptx_path = OUT / "browser-pptx.pptx"
    if pptx_path.exists():
        deck = Presentation(pptx_path)
        check("browser pptx: edited title", deck.slides[0].shapes.title.text, "Q4 Review — browser")
        check("browser pptx: other slide intact", deck.slides[1].shapes.title.text, "Highlights")
        check("browser pptx: layouts intact", len(deck.slide_layouts), 11)
    else:
        print(f"  skip {pptx_path.name} (run the browser suite first)")

    html_path = OUT / "browser-code.html"
    if html_path.exists():
        markup = html_path.read_text(encoding="utf-8")
        check("browser html: typed marker present", "<!-- edited -->" in markup, True)
        check("browser html: original markup intact", "Quarterly Report" in markup, True)
    else:
        print(f"  skip {html_path.name} (run the browser suite first)")


def main() -> int:
    if not OUT.exists():
        print(f"missing {OUT} — run the vitest real-fixtures test first", file=sys.stderr)
        return 2

    verify_xlsx()
    verify_pptx()
    verify_browser_outputs()

    print()
    if failures:
        print(f"{len(failures)} check(s) FAILED:")
        for failure in failures:
            print(f"  - {failure}")
        return 1

    print("all reference-library checks passed")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
